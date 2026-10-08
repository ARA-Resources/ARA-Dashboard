"""
Build a synthetic .xlsm-shaped fixture (saved as .xlsx — openpyxl can't
fabricate a real VBA project from scratch, so macro/vbaProject preservation
is NOT exercised by this fixture; see the integration test's own note on
that limitation) for testing the Posted-button local writer without any
real Drive credentials or real workbook.

Tabs mirror the real workbook's shape: Team Guidelines, P-Roles, Master
Sheet, New Sheet, Posted Sheet, Allocation Sheet, Instructions. Posted Sheet
has a header row, an AutoFilter already spanning A:D (deliberately NOT A:C,
to prove the writer preserves whatever column span already exists rather
than assuming one), a blank row, a title line, a clean canonical row, and a
JR-ID-but-no-city row. Master Sheet has a couple of JR rows. P-Roles has a
simple formula referencing Master Sheet, to prove it survives untouched.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font

out_path = sys.argv[1]

wb = Workbook()
wb.remove(wb.active)

guidelines = wb.create_sheet("Team Guidelines")
guidelines["A1"] = "Please follow the process."

p_roles = wb.create_sheet("P-Roles")
p_roles["A1"] = "Pivot"
p_roles["B1"] = "=COUNTIF('Master Sheet'!B2:B3,\"Yes\")"

master = wb.create_sheet("Master Sheet")
master["A1"] = "Job Requisition ID"
master["B1"] = "Posted"
master["A2"] = "ATCI-7001-S1"
master["B2"] = "-"
master["A3"] = "ATCI-9999-S9"
master["B3"] = "Yes"

new_sheet = wb.create_sheet("New Sheet")
new_sheet["A1"] = "Date"
new_sheet["B1"] = "Job Requisition ID"

posted = wb.create_sheet("Posted Sheet")
posted["A1"] = "Job Requisition"
posted["B1"] = "Job Requisition ID"
posted["C1"] = "Demand"
posted["A2"] = None  # blank row
posted["A3"] = "Custom Software Engineer"
posted["A3"].font = Font(name="Calibri", size=14)  # the "one row looks bigger" bug
posted["A4"] = "ATCI-7001-S1 | Posting Date: 01/01/2026 | Pune"
posted["A4"].font = Font(name="Calibri", size=11)
posted["A5"] = None  # blank row
posted["A6"] = "ATCI-7002-S1 | Posting Date: 02/02/2026"  # JR ID + date, no city
posted["A6"].font = Font(name="Calibri", size=11)
posted.auto_filter.ref = "A1:D6"

allocation = wb.create_sheet("Allocation Sheet")
allocation["A1"] = "Allocation"

instructions = wb.create_sheet("Instructions")
instructions["A1"] = "Read me first."

wb.save(out_path)
print("fixture built:", out_path)
