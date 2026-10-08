"""
Fixture matching Step 18's REAL hardcoded assumptions (lateral-posted-
sheet-processor.ts): Master Sheet Job Requisition ID at column B (index 2,
hardcoded), Posted written at column M (index 13, MASTER_POSTED_COLUMN_M).
Our own writer resolves both by header name, so this single layout is
correct for both at once — needed for the Step 18 round-trip comparison.
"""
import sys
from openpyxl import Workbook

out_path = sys.argv[1]

wb = Workbook()
wb.remove(wb.active)

master = wb.create_sheet("Master Sheet")
headers = ["Date", "Job Requisition ID", "Priority", "Job Description", "Skill Categorization",
           "Primary Skills", "Job Management Level", "Primary Location", "Market Map", "POC",
           "Job Status", "Filler", "Posted"]
for i, h in enumerate(headers, start=1):
    master.cell(1, i).value = h
master.cell(2, 2).value = "ATCI-7001-S1"
master.cell(2, 13).value = "-"
master.cell(3, 2).value = "ATCI-7002-S1"
master.cell(3, 13).value = "-"
master.cell(4, 2).value = "ATCI-9999-S9"
master.cell(4, 13).value = "Yes"

new_sheet = wb.create_sheet("New Sheet")
new_sheet["A1"] = "Date"
new_sheet["B1"] = "Job Requisition ID"

posted = wb.create_sheet("Posted Sheet")
posted["A1"] = "Job Requisition"
posted["B1"] = "Job Requisition ID"
posted["C1"] = "Demand"
posted["A2"] = "ATCI-7001-S1 | Posting Date: 01/01/2026 | Pune"
posted["A3"] = "ATCI-7002-S1 | Posting Date: 02/02/2026"  # JR + date, no city

wb.save(out_path)
print("step18 fixture built:", out_path)
