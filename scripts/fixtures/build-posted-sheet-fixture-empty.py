"""Zero-JR-ID fixture (blank row + a title line only) for the zero-IDs guard test."""
import sys
from openpyxl import Workbook

out_path = sys.argv[1]

wb = Workbook()
wb.remove(wb.active)

p_roles = wb.create_sheet("P-Roles")
p_roles["A1"] = "Pivot"

master = wb.create_sheet("Master Sheet")
master["A1"] = "Job Requisition ID"
master["B1"] = "Posted"
master["A2"] = "ATCI-7001-S1"
master["B2"] = "-"

posted = wb.create_sheet("Posted Sheet")
posted["A1"] = "Job Requisition"
posted["B1"] = "Job Requisition ID"
posted["C1"] = "Demand"
posted["A2"] = None
posted["A3"] = "Custom Software Engineer"
posted.auto_filter.ref = "A1:D3"

wb.save(out_path)
print("empty fixture built:", out_path)
